"""CV ingestion & structured-profile extraction.

Two paths:
1. ``extract_cv_text`` — pull raw text out of a PDF/TXT/MD upload.
2. ``parse_profile`` — deterministic extraction of a structured candidate
   profile (name, contact, experience, education, skills) from raw CV text.

``parse_profile`` is deliberately deterministic so the pipeline is reproducible
offline. When ``LLM_BACKEND=bedrock`` and ``USE_LLM_PARSE=1``, a Bedrock pass can
enrich the profile for scanned/image CVs; on failure it falls back gracefully.
"""
from __future__ import annotations

import io
import re
from typing import Optional

from app.models.schemas import CandidateProfile
from app.services.llm import get_llm

# Known-skill dictionary used to tag skills from free-form CV text. Keys are
# canonical labels; values are the aliases that signal the skill.
SKILL_ALIASES: dict[str, list[str]] = {
    "Python": ["python"],
    "REST APIs": ["rest api", "restapis", "restful", "rest"],
    "FastAPI": ["fastapi", "fast api"],
    "Docker": ["docker"],
    "Google Cloud": ["google cloud", "gcp", "google cloud platform"],
    "Kubernetes": ["kubernetes", "k8s"],
    "SQL": ["sql"],
    "PostgreSQL": ["postgresql", "postgres"],
    "MySQL": ["mysql"],
    "Redis": ["redis"],
    "AWS": ["aws", "amazon web services"],
    "Azure": ["azure"],
    "JavaScript": ["javascript", " js "],
    "TypeScript": ["typescript", "ts "],
    "React": ["react"],
    "Node.js": ["node.js", "nodejs", "node "],
    "Go": ["golang", "go lang"],
    "Java": ["java"],
    "C++": ["c++", "cpp"],
    "Machine Learning": ["machine learning", " ml "],
}

_EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
_PHONE_RE = re.compile(r"(\+?\d[\d\s\-()]{8,}\d)")
_YEARS_RE = re.compile(r"(\d+(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)", re.IGNORECASE)
_EXPERIENCE_LABEL_RE = re.compile(r"experience\s*:?\s*(\d+(?:\.\d+)?)", re.IGNORECASE)


def _as_float(value) -> Optional[float]:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def extract_cv_text(file_bytes: bytes, filename: str = "") -> str:
    """Return raw text from an uploaded CV file (PDF, DOCX, TXT, or MD)."""
    name = (filename or "").lower()

    if name.endswith(".pdf"):
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(file_bytes))
            return "\n".join((page.extract_text() or "") for page in reader.pages)
        except Exception:
            return ""  # unparseable -> caller escalates

    if name.endswith(".docx"):
        try:
            from docx import Document

            doc = Document(io.BytesIO(file_bytes))
            parts: list[str] = []
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    parts.append(paragraph.text)
            # Include text inside tables (common in CV layouts).
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            parts.append(cell.text)
            return "\n".join(parts)
        except Exception:
            return ""  # unparseable -> caller escalates

    # TXT / MD / anything else: assume UTF-8 text
    for enc in ("utf-8", "latin-1"):
        try:
            return file_bytes.decode(enc)
        except UnicodeDecodeError:
            continue
    return ""


def parse_profile(cv_text: str, name: str = "", email: str = "", phone: str = "") -> CandidateProfile:
    """Deterministically extract a structured profile from CV text.

    ``name`` / ``email`` / ``phone`` (when supplied from the application form)
    take precedence over anything found in the document.
    """
    text = cv_text or ""
    text_lower = text.lower()

    extracted_email = email or (_EMAIL_RE.search(text).group(0) if _EMAIL_RE.search(text) else "")
    extracted_phone = phone or (_PHONE_RE.search(text).group(0).strip() if _PHONE_RE.search(text) else "")

    experience_years = _extract_experience_years(text)
    education = _extract_education(text)
    skills = _extract_skills(text_lower)

    if not name:
        name = _extract_name(text)

    return CandidateProfile(
        name=name.strip() or "Unknown Candidate",
        email=extracted_email,
        phone=extracted_phone,
        experience_years=experience_years,
        education=education,
        skills=skills,
    )


def _extract_experience_years(text: str) -> Optional[float]:
    for m in _YEARS_RE.finditer(text):
        return float(m.group(1))
    for m in _EXPERIENCE_LABEL_RE.finditer(text):
        return float(m.group(1))
    return None


def _extract_education(text: str) -> list[str]:
    edu_keywords = ("bachelor", "b.sc", "b. sc", "bs ", "b.s", "master", "m.sc", "m. sc",
                    "ms ", "m.s", "phd", "ph.d", "degree", "university", "b.tech", "m.tech")
    found: list[str] = []
    for line in text.splitlines():
        low = line.lower().strip()
        if any(k in low for k in edu_keywords) and len(line.strip()) < 120:
            found.append(line.strip())
    return found[:4]


def _extract_skills(text_lower: str) -> list[str]:
    skills: list[str] = []
    # Explicit "Skills:" block first
    skills_line = re.search(r"skills\s*:?\s*(.+)", text_lower)
    if skills_line:
        block = re.split(r"[\n,;|•]", skills_line.group(1))
        for tok in block:
            tok = tok.strip().rstrip(".")
            if 1 < len(tok) < 40:
                skills.append(tok.title())
    # Tag against the known dictionary across the whole document
    for label, aliases in SKILL_ALIASES.items():
        if any(a in text_lower for a in aliases):
            if label.lower() not in [s.lower() for s in skills]:
                skills.append(label)
    return skills


def _extract_name(text: str) -> str:
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    for line in lines[:6]:
        low = line.lower()
        if "@" in line or _PHONE_RE.search(line):
            continue
        if any(w in low for w in ("curriculum vitae", "resume", "cv", "name:")):
            continue
        if len(line.split()) <= 5 and len(line) < 60:
            return line
    return "Unknown Candidate"


def parse_profile_with_llm(cv_text: str) -> CandidateProfile:
    """Optional Bedrock-assisted parse (used for scanned/image CVs)."""
    try:
        llm = get_llm()
        prompt = (
            "Extract the following fields from this CV as JSON: "
            'name, email, phone, experience_years (number), education (array), '
            'skills (array). Return only JSON.\n\nCV:\n"""\n' + cv_text[:6000] + '\n"""'
        )
        data = llm.generate_json(prompt)
        return CandidateProfile(
            name=data.get("name", "Unknown Candidate"),
            email=data.get("email", ""),
            phone=data.get("phone", ""),
            experience_years=_as_float(data.get("experience_years")),
            education=data.get("education", []) or [],
            skills=data.get("skills", []) or [],
        )
    except Exception:
        return parse_profile(cv_text)
